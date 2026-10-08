"""Extract and rasterize images embedded in PPTX (e.g. WMF) when MarkItDown cannot inline them."""

from __future__ import annotations

import base64
import io
import logging
import os
import re
import subprocess
import tempfile
import uuid
import zipfile
from collections import deque
from typing import Dict, List, Tuple

logger = logging.getLogger(__name__)

_MARKDOWN_IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
_RASTER_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
_VECTOR_EXT = {".wmf", ".emf", ".svg"}


def _find_convert() -> str | None:
    for path in ("/usr/bin/convert", "/usr/local/bin/convert"):
        if os.path.isfile(path):
            return path
    try:
        result = subprocess.run(
            ["which", "convert"], capture_output=True, text=True, check=False
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()
    except OSError:
        pass
    return None


def _rasterize_with_imagemagick(data: bytes, suffix: str) -> bytes | None:
    convert = _find_convert()
    if not convert:
        return None
    with tempfile.TemporaryDirectory() as temp_dir:
        src = os.path.join(temp_dir, f"input{suffix}")
        dst = os.path.join(temp_dir, "output.png")
        with open(src, "wb") as handle:
            handle.write(data)
        try:
            result = subprocess.run(
                [convert, src, dst],
                capture_output=True,
                timeout=60,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            logger.warning("ImageMagick convert failed: %s", exc)
            return None
        if result.returncode != 0 or not os.path.isfile(dst):
            stderr = (result.stderr or b"").decode("utf-8", errors="ignore")
            logger.warning("ImageMagick convert exit %s: %s", result.returncode, stderr)
            return None
        with open(dst, "rb") as handle:
            return handle.read()


def _rasterize_with_pillow(data: bytes) -> bytes | None:
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        img = Image.open(io.BytesIO(data))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        out = io.BytesIO()
        img.save(out, format="PNG")
        return out.getvalue()
    except Exception as exc:
        logger.debug("Pillow could not open media bytes: %s", exc)
        return None


def rasterize_media_bytes(name: str, data: bytes) -> bytes | None:
    ext = os.path.splitext(name)[1].lower()
    if ext in _RASTER_EXT:
        png = _rasterize_with_pillow(data)
        if png:
            return png
    if ext in _VECTOR_EXT or ext in _RASTER_EXT:
        return _rasterize_with_imagemagick(data, ext or ".bin")
    return _rasterize_with_imagemagick(data, ext or ".bin")


def list_pptx_media(pptx_bytes: bytes) -> List[Tuple[str, bytes]]:
    """Return (zip path, raw bytes) for each file under ppt/media/, in archive order."""
    items: List[Tuple[str, bytes]] = []
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as archive:
        for name in archive.namelist():
            if not name.startswith("ppt/media/"):
                continue
            base = os.path.basename(name)
            if not base or base.startswith("."):
                continue
            items.append((name, archive.read(name)))
    return items


def extract_pptx_media_rasterized(pptx_bytes: bytes) -> List[bytes]:
    """Rasterize all ppt/media assets to PNG bytes, skipping failures."""
    rasterized: List[bytes] = []
    for path, raw in list_pptx_media(pptx_bytes):
        png = rasterize_media_bytes(os.path.basename(path), raw)
        if png:
            rasterized.append(png)
            logger.info(
                "Rasterized pptx media %s (%d -> %d bytes)", path, len(raw), len(png)
            )
        else:
            logger.warning("Failed to rasterize pptx media %s", path)
    return rasterized


def _is_unresolved_image_ref(url: str) -> bool:
    if not url or url.startswith("data:") or url.startswith("images/"):
        return False
    if url.startswith(("http://", "https://")):
        return False
    return True


def _picture_sources(pptx_bytes: bytes) -> dict[str, deque[tuple[str, bytes] | None]]:
    """Resolve MarkItDown picture placeholders in slide/shape order."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.exc import PythonPptxError

    sources: dict[str, deque[tuple[str, bytes] | None]] = {}

    def visit(shapes) -> None:
        # Match MarkItDown, which emits shapes (and group members) sorted by
        # position rather than in z-order.
        for shape in sorted(
            shapes,
            key=lambda x: (
                float("-inf") if not x.top else x.top,
                float("-inf") if not x.left else x.left,
            ),
        ):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes)
                continue
            is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (
                shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
                and hasattr(shape, "image")
            )
            if not is_picture:
                continue

            # MarkItDown's no-data-URI mode uses the shape name, not the media
            # part name. Names may repeat across slides or normalize equally.
            placeholder = re.sub(r"\W", "", shape.name) + ".jpg"
            source = None
            try:
                blip = shape._element.blipFill.blip
                relationship = blip.rEmbed if blip is not None else None
                if relationship:
                    part = shape.part.related_part(relationship)
                    # Read the part directly: shape.image.content_type can
                    # require a decoder for the same vector format that
                    # triggered fallback.
                    source = (str(part.partname), part.blob)
            except (AttributeError, KeyError, ValueError, PythonPptxError) as exc:
                # A missing blipFill or blip, a dangling or external r:embed,
                # or a media part dropped at load time must only lose this
                # picture, not the whole document. None keeps the queue
                # position.
                logger.warning(
                    "Skipping unresolvable pptx picture %s: %s", shape.name, exc
                )
            sources.setdefault(placeholder, deque()).append(source)

    presentation = Presentation(io.BytesIO(pptx_bytes))
    for slide in presentation.slides:
        visit(slide.shapes)
    return sources


def attach_pptx_media_to_markdown(
    markdown: str, pptx_bytes: bytes
) -> Tuple[str, Dict[str, str]]:
    """Attach the media referenced by MarkItDown's fallback picture placeholders."""
    sources = _picture_sources(pptx_bytes)
    if not sources:
        return markdown, {}

    images: Dict[str, str] = {}
    converted: dict[str, bytes | None] = {}

    def repl(match: re.Match[str]) -> str:
        alt, url = match.group(1), match.group(2)
        if not _is_unresolved_image_ref(url):
            return match.group(0)
        pending = sources.get(url)
        if not pending:
            return match.group(0)
        source = pending.popleft()
        if source is None:
            return match.group(0)
        name, raw = source
        if name not in converted:
            converted[name] = rasterize_media_bytes(name, raw)
            if converted[name] is None:
                logger.warning("Failed to rasterize pptx media %s", name)
        png = converted[name]
        if not png:
            return match.group(0)
        ref = f"images/{uuid.uuid4()}.png"
        images[ref] = base64.b64encode(png).decode()
        return f"![{alt}]({ref})"

    return _MARKDOWN_IMAGE.sub(repl, markdown), images


def markdown_needs_pptx_media_attach(markdown: str) -> bool:
    return any(
        _is_unresolved_image_ref(m.group(2)) for m in _MARKDOWN_IMAGE.finditer(markdown)
    )
