import base64
import io
import re
import shutil
import unittest
import zipfile
from unittest.mock import patch

from lxml import etree
from markitdown import MarkItDown
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from docreader.parser import pptx_media
from docreader.parser.markitdown_parser import MarkitdownParser, StdMarkitdownParser

RED = (220, 30, 40)
BLUE = (30, 80, 220)
GREEN = (20, 170, 70)
IMAGE = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")


def png(color):
    output = io.BytesIO()
    Image.new("RGB", (16, 16), color).save(output, format="PNG")
    return output.getvalue()


def deck(colors, names=None, grouped=False, reverse_slides=False):
    presentation = Presentation()
    for index, color in enumerate(colors):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shapes = slide.shapes.add_group_shape().shapes if grouped else slide.shapes
        picture = shapes.add_picture(io.BytesIO(png(color)), Inches(1), Inches(1))
        picture.name = names[index] if names else f"picture{index}"
        picture._element._nvXxPr.cNvPr.set("descr", f"position{index}")
    if reverse_slides:
        ids = presentation.slides._sldIdLst
        for child in list(ids):
            ids.remove(child)
            ids.insert(0, child)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def rewrite_archive(data, *, reverse_media=False, unused=False, corrupt_first=False):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    media = [name for name in files if name.startswith("ppt/media/")]
    if corrupt_first:
        files[media[0]] = b"not an image"
    if reverse_media:
        media.reverse()
    if unused:
        name = "ppt/media/unused.png"
        files[name] = png(GREEN)
        media.insert(0, name)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in files.items():
            if name not in media:
                archive.writestr(name, blob)
        for name in media:
            archive.writestr(name, files[name])
    return output.getvalue()


def fallback(data):
    markdown = (
        MarkItDown()
        .convert(io.BytesIO(data), file_extension=".pptx", keep_data_uris=False)
        .text_content
    )
    return pptx_media.attach_pptx_media_to_markdown(markdown, data)


def image_colors(markdown, images):
    colors = []
    for target in IMAGE.findall(markdown):
        if target not in images:
            colors.append(None)
            continue
        raw = base64.b64decode(images[target])
        with Image.open(io.BytesIO(raw)) as image:
            colors.append(image.convert("RGB").getpixel((8, 8)))
    return colors


class TestPptxMediaAssociation(unittest.TestCase):
    def test_document_order_control(self):
        markdown, images = fallback(deck([RED, BLUE]))
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_zip_entry_order_does_not_change_which_picture_is_attached(self):
        data = rewrite_archive(deck([RED, BLUE]), reverse_media=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_unused_media_does_not_replace_a_visible_picture(self):
        data = rewrite_archive(deck([RED, BLUE]), unused=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_shared_image_is_attached_at_every_placement(self):
        markdown, images = fallback(deck([RED, RED, BLUE]))
        self.assertEqual(image_colors(markdown, images), [RED, RED, BLUE])

    def test_failed_conversion_does_not_take_the_next_picture(self):
        data = rewrite_archive(deck([RED, BLUE]), corrupt_first=True)
        # A missing converter is a supported deployment condition. The bad
        # PNG still exercises Pillow's actual decode failure.
        with patch.object(pptx_media, "_find_convert", return_value=None):
            markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [None, BLUE])

    def test_duplicate_shape_names_keep_each_placement(self):
        data = deck([RED, BLUE], names=["Picture 1", "Picture 1"])
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_names_that_share_a_placeholder_keep_each_placement(self):
        data = rewrite_archive(
            deck([RED, BLUE], names=["Picture-1", "Picture 1"]), reverse_media=True
        )
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_presentation_order_takes_precedence_over_slide_filenames(self):
        markdown, images = fallback(deck([RED, BLUE], reverse_slides=True))
        self.assertEqual(image_colors(markdown, images), [BLUE, RED])

    def test_same_slide_duplicate_names_follow_position_order(self):
        # MarkItDown orders shapes by (top, left), not z-order.
        for grouped in (False, True):
            with self.subTest(grouped=grouped):
                presentation = Presentation()
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                shapes = (
                    slide.shapes.add_group_shape().shapes if grouped else slide.shapes
                )
                for color, top in ((RED, 4), (BLUE, 1)):
                    picture = shapes.add_picture(
                        io.BytesIO(png(color)), Inches(1), Inches(top)
                    )
                    picture.name = "Picture 1"
                output = io.BytesIO()
                presentation.save(output)
                markdown, images = fallback(output.getvalue())
                self.assertEqual(image_colors(markdown, images), [BLUE, RED])

    def test_grouped_pictures_keep_their_source(self):
        data = rewrite_archive(deck([RED, BLUE], grouped=True), reverse_media=True)
        markdown, images = fallback(data)
        self.assertEqual(image_colors(markdown, images), [RED, BLUE])

    def test_external_picture_does_not_consume_the_next_embedded_image(self):
        presentation = Presentation(io.BytesIO(deck([RED, BLUE])))
        slide = presentation.slides[0]
        picture = slide.shapes[0]
        blip = picture._element.blipFill.blip
        slide.part.drop_rel(blip.rEmbed)
        del blip.attrib[qn("r:embed")]
        relationship = slide.part.relate_to(
            "https://example.invalid/external.png",
            RELATIONSHIP_TYPE.IMAGE,
            is_external=True,
        )
        blip.set(qn("r:link"), relationship)
        output = io.BytesIO()
        presentation.save(output)
        markdown, images = fallback(output.getvalue())
        self.assertEqual(image_colors(markdown, images), [None, BLUE])

    def test_unrelated_links_and_unknown_placeholders_stay_unchanged(self):
        data = deck([RED])
        original = (
            MarkItDown()
            .convert(io.BytesIO(data), file_extension=".pptx", keep_data_uris=False)
            .text_content
        )
        prefix = (
            "![web](https://example.invalid/image.png)\n"
            "![inline](data:image/png;base64,eA==)\n"
            "![unknown](not-a-picture.jpg)\n"
        )
        markdown, images = pptx_media.attach_pptx_media_to_markdown(
            prefix + original, data
        )
        self.assertTrue(markdown.startswith(prefix))
        self.assertEqual(image_colors(markdown, images)[-1], RED)


def deck_with_broken_picture(mutate):
    """One slide: text, a good RED picture, then a BLUE picture `mutate` breaks."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(0.1), Inches(4), Inches(0.5))
    textbox.text_frame.text = "SLIDE TEXT OK"
    good = slide.shapes.add_picture(io.BytesIO(png(RED)), Inches(1), Inches(1))
    good.name = "good"
    bad = slide.shapes.add_picture(io.BytesIO(png(BLUE)), Inches(1), Inches(3))
    bad.name = "bad"
    media = mutate(slide, bad)
    output = io.BytesIO()
    presentation.save(output)
    data = output.getvalue()
    if media is None:
        return data
    # Drop the picture's media part from the package: python-pptx then drops
    # the relationship at load time while the slide XML still points at it.
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    del files[media]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in files.items():
            archive.writestr(name, blob)
    return output.getvalue()


def remove_blip(slide, picture):
    blip = picture._element.blipFill.blip
    blip.getparent().remove(blip)


def remove_blip_fill(slide, picture):
    blip_fill = picture._element.blipFill
    blip_fill.getparent().remove(blip_fill)


def dangling_embed(slide, picture):
    picture._element.blipFill.blip.set(qn("r:embed"), "rId999")


def external_embed(slide, picture):
    relationship = slide.part.relate_to(
        "https://example.invalid/external.png",
        RELATIONSHIP_TYPE.IMAGE,
        is_external=True,
    )
    picture._element.blipFill.blip.set(qn("r:embed"), relationship)


def missing_media_part(slide, picture):
    part = slide.part.related_part(picture._element.blipFill.blip.rEmbed)
    return str(part.partname).lstrip("/")


class TestPptxMediaBrokenPictureRelationships(unittest.TestCase):
    def test_broken_picture_keeps_the_document_and_other_pictures(self):
        for mutate in (
            remove_blip,
            remove_blip_fill,
            dangling_embed,
            external_embed,
            missing_media_part,
        ):
            with self.subTest(mutate.__name__):
                data = deck_with_broken_picture(mutate)
                document = StdMarkitdownParser(
                    file_name="deck.pptx", file_type="pptx"
                ).parse_into_text(data)
                self.assertIn("SLIDE TEXT OK", document.content)
                self.assertEqual(
                    image_colors(document.content, document.images), [RED, None]
                )
                self.assertIn("(bad.jpg)", document.content)

                parsed = MarkitdownParser(
                    file_name="deck.pptx", file_type="pptx"
                ).parse(data)
                self.assertIn("SLIDE TEXT OK", parsed.content)


class TestPptxMediaRealFallback(unittest.TestCase):
    @unittest.skipUnless(shutil.which("convert"), "ImageMagick is not installed")
    def test_svg_fallback_preserves_picture_order(self):
        # SVG is valid Office media, but Pillow cannot identify it. That makes
        # pinned MarkItDown 0.1.3 naturally enter the actual fallback path.
        data = deck([RED, BLUE])
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        original = "ppt/media/image2.png"
        replacement = "ppt/media/image2.svg"
        files[replacement] = (
            b'<svg xmlns="http://www.w3.org/2000/svg" width="16" height="16">'
            b'<rect width="16" height="16" fill="#1e50dc"/></svg>'
        )
        del files[original]
        for name, blob in list(files.items()):
            if name.endswith(".rels"):
                files[name] = blob.replace(b"image2.png", b"image2.svg")
        types = etree.fromstring(files["[Content_Types].xml"])
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
            Extension="svg",
            ContentType="image/svg+xml",
        )
        files["[Content_Types].xml"] = etree.tostring(types)
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(replacement, files.pop(replacement))
            for name, blob in files.items():
                archive.writestr(name, blob)

        with self.assertLogs("docreader.parser.markitdown_parser", level="WARNING"):
            document = MarkitdownParser(file_type="pptx").parse(output.getvalue())
        self.assertEqual(image_colors(document.content, document.images), [RED, BLUE])


if __name__ == "__main__":
    unittest.main()
